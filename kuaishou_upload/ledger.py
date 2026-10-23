from __future__ import annotations

import os
import shutil
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .database import DONE
from .state import exclusive_lock

PUBLISHED = {"published", "submitted", "exists", "已存在-跳过", "已发布", "已上传"}
HEADERS = ["文件名", "本地路径", "发布状态", "发布时间", "审核状态", "最近检查时间", "备注",
           "文件哈希", "剧名", "集数", "作品ID", "集数ID", "任务ID", "平台", "账号"]


def row_scope(row):
    # Unscoped historical rows belong only to the original account.
    return str(row.get('平台') or 'kuaishou'), str(row.get('账号') or 'default')


def normalized(path):
    return os.path.normcase(os.path.normpath(str(path)))


class Ledger:
    def __init__(self, path: Path):
        self.path = path

    def read(self):
        if not self.path.exists():
            return []
        book = load_workbook(self.path, read_only=True, data_only=False)
        try:
            if "发布记录" not in book.sheetnames:
                raise ValueError("现有台账缺少“发布记录”工作表，停止以免丢失历史查重信息")
            rows = list(book["发布记录"].iter_rows(values_only=True))
            names = list(rows[0]) if rows else []
            if len([n for n in names if n]) != len({n for n in names if n}):
                raise ValueError("台账含重复列名，请先核对")
            return [{**{name: row[i] if i < len(row) else None for i, name in enumerate(names) if name},
                     "_row": index} for index, row in enumerate(rows[1:], 2)]
        finally:
            book.close()

    @staticmethod
    def match(rows, video: Path, sha256: str, *, unique_name: bool, platform='kuaishou', account_id='default'):
        matches = []
        for row in rows:
            if row_scope(row) != (platform, account_id):
                continue
            old_hash = str(row.get("文件哈希") or "").lower()
            by_path = row.get("本地路径") and normalized(row["本地路径"]) == normalized(video)
            by_name = unique_name and str(row.get("文件名") or "").casefold() == video.name.casefold()
            if old_hash:
                if old_hash == sha256:
                    matches.append(row)
                elif by_path or by_name:
                    raise ValueError("台账同名/同路径视频的哈希不同，需核对是否替换过文件")
            elif by_path or by_name:
                matches.append(row)
        if len(matches) > 1:
            raise ValueError("台账存在多条匹配记录，请先核对，禁止自动重传")
        return matches[0] if matches else None

    def write(self, record):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with exclusive_lock(self.path.with_suffix(".xlsx.lock")):
            if self.path.exists():
                backup = self.path.with_name(self.path.stem + ".before-agent.xlsx")
                if not backup.exists():
                    shutil.copy2(self.path, backup)
                book = load_workbook(self.path)
                if "发布记录" not in book.sheetnames:
                    book.close()
                    raise ValueError("现有台账缺少发布记录工作表")
                sheet = book["发布记录"]
            else:
                book = Workbook()
                sheet = book.active
                sheet.title = "发布记录"
            try:
                headers = {c.value: c.column for c in sheet[1] if c.value}
                for name in HEADERS:
                    if name not in headers:
                        headers[name] = max(headers.values(), default=0) + 1
                        sheet.cell(1, headers[name], name)
                matches = [r for r in range(2, sheet.max_row + 1)
                           if str(sheet.cell(r, headers["文件哈希"]).value or "").lower() == record["sha256"]
                           and row_scope({k: sheet.cell(r, headers[k]).value for k in ('平台', '账号')})
                           == (record.get('platform', 'kuaishou'), record.get('account_id', 'default'))]
                if not matches:
                    old_row = record["details"].get("ledgerRow")
                    if old_row and 2 <= old_row <= sheet.max_row:
                        same_file = (sheet.cell(old_row, headers["文件名"]).value == Path(record["path"]).name
                                     or normalized(sheet.cell(old_row, headers["本地路径"]).value or "") == normalized(record["path"]))
                        old_hash = sheet.cell(old_row, headers["文件哈希"]).value
                        scope = row_scope({k: sheet.cell(old_row, headers[k]).value for k in ('平台', '账号')})
                        if same_file and scope == (record.get('platform', 'kuaishou'), record.get('account_id', 'default')) and (not old_hash or str(old_hash).lower() == record["sha256"]):
                            matches = [old_row]
                if len(matches) > 1:
                    raise ValueError("同一视频在台账中有多行，拒绝覆盖")
                row = matches[0] if matches else max(sheet.max_row + 1, 2)
                details = record["details"]
                values = {"文件名": Path(record["path"]).name, "本地路径": record["path"],
                          "发布状态": "published" if record["status"] in DONE else record["status"],
                          "最近检查时间": record["updated_at"], "文件哈希": record["sha256"],
                          "任务ID": record['details'].get('jobId', record["sha256"][:24]),
                          "平台": record.get('platform', 'kuaishou'), "账号": record.get('account_id', 'default'),
                          "剧名": record["series_title"], "集数": record["episode"],
                          "作品ID": record["content_id"], "集数ID": record["episode_id"]}
                if details.get("submittedAt"):
                    values["发布时间"] = details["submittedAt"]
                if details.get("reviewStatus"):
                    values["审核状态"] = details["reviewStatus"]
                previous = str(sheet.cell(row, headers["备注"]).value or "").split("\n[上传工具]", 1)[0]
                values["备注"] = previous + "\n[上传工具] " + record["message"]
                for name, value in values.items():
                    if value is None:
                        continue
                    cell = sheet.cell(row, headers[name], value)
                    if isinstance(value, str):
                        cell.data_type = "s"
                fd, filename = tempfile.mkstemp(dir=self.path.parent, suffix=".xlsx")
                os.close(fd)
                try:
                    book.save(filename)
                    os.replace(filename, self.path)
                finally:
                    if os.path.exists(filename):
                        os.unlink(filename)
            finally:
                book.close()


def sync_ledger(database, ledger):
    failures = []
    for record in database.records(dirty_only=True):
        try:
            ledger.write(record)
            database.update(record["sha256"], ledger_dirty=0)
        except Exception as exc:
            failures.append(str(exc))
            break  # Excel 被占用时保留待同步标记，下次运行再同步。
    return failures
