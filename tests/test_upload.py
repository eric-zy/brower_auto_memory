import asyncio
import copy
import json
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from openpyxl import Workbook, load_workbook

from kuaishou_upload.browser import Decision, PlatformBrowser, safe_control
from kuaishou_upload.config import Settings
from kuaishou_upload.database import UploadDatabase
from kuaishou_upload.ledger import Ledger, sync_ledger
from kuaishou_upload.scan import scan_directory
from kuaishou_upload.schema import file_hash, load_job
from kuaishou_upload.state import atomic_json, exclusive_lock
from kuaishou_upload.tools import JobParams, FormParams, InspectParams, ScanParams, UploadParams, UploadTask, build_tools
from support import make_materials


class FakeBrowser:
    def __init__(self, job):
        self.job = job
        self.page = SimpleNamespace(get_url=AsyncMock(return_value="https://kdj.kuaishou.com/home/content/content-management"))
        self.form_verified = True
        self.filled = {}
        self.expected_card_count = 1
        self.uploaded = set()
        self.checkpoint = None
        self.snap = {"url": "https://kdj.kuaishou.com/home/content/content-management", "text": "内容管理 暂无数据",
                     "controls": [{"ref": "k1", "value": job.manifest.seriesTitle, "type": "text", "tag": "input", "text": "", "label": "搜索剧名"}],
                     "rows": [], "uploadRows": [], "dialogs": []}
        self.act = AsyncMock()
        self.verify_prepared = AsyncMock()
        self.submit = AsyncMock(return_value={"contentId": "C1", "episodeId": "E1", "reviewStatus": "待审核"})
        self.reconcile = AsyncMock(return_value=None)
        self.open_detail = AsyncMock()
        self.screenshot = AsyncMock(return_value=None)
        self.record_readback = lambda *args: None

    async def snapshot(self):
        return copy.deepcopy(self.snap)

    def find_episode(self, snap):
        return None

    def video_complete(self, snap, name):
        return PlatformBrowser.video_complete(snap, name)


class UploadTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.base = Path(self.tmp.name)
        self.directory = self.base / "videos"
        self.sidecar = make_materials(self.directory)
        self.settings = Settings(self.directory, self.base / "state", self.directory / "发布记录.xlsx", stable_seconds=0)
        self.db = UploadDatabase(self.settings.root / "uploads.sqlite3")

    def tearDown(self):
        self.db.close()
        self.tmp.cleanup()

    def job(self):
        return load_job(self.sidecar, allow_test_prepare=True)

    def scan(self):
        return scan_directory(self.settings, self.db)

    def context(self, status="pending", **details):
        job = self.job()
        self.scan()
        sha = job.manifest.mediaInfo.sha256
        self.db.update(sha, status=status, content_id="C1", details=details)
        return UploadTask(job, sha, self.settings, self.db, {"counts": {"pending": 1}}, browser=FakeBrowser(job))

    def modify(self, function):
        data = json.loads(self.sidecar.read_text(encoding="utf-8"))
        function(data)
        atomic_json(self.sidecar, data)

    def legacy(self, rows, headers=None):
        book = Workbook()
        sheet = book.active
        sheet.title = "发布记录"
        sheet.append(headers or ["文件名", "本地路径", "发布状态", "审核状态", "备注"])
        for row in rows:
            sheet.append(row)
        book.create_sheet("保留数据")["A1"] = "keep"
        book.save(self.settings.ledger)
        book.close()

    def test_scan_requires_no_model_and_queues_valid_job(self):
        with patch.dict(os.environ, {"OPENAI_API_KEY": ""}):
            result = self.scan()
        self.assertEqual(len(result.jobs), 1)
        self.assertEqual(result.summary()["counts"], {"pending": 1})

    def test_scanner_excludes_testdata_and_temporary_files(self):
        make_materials(self.directory / "testdata", name="test")
        (self.directory / "not-video.mp4.downloading").write_bytes(b"x")
        result = self.scan()
        self.assertEqual(len(result.items), 1)

    def test_downloading_sibling_defers_video(self):
        (self.directory / "demo.mp4.downloading").write_bytes(b"x")
        result = self.scan()
        self.assertEqual(result.items[0]["status"], "deferred")
        self.assertEqual(self.db.records(), [])

    def test_recent_file_is_not_uploaded(self):
        self.settings.stable_seconds = 120
        self.assertEqual(self.scan().items[0]["status"], "deferred")

    def test_legacy_published_is_skipped_before_missing_material_check(self):
        self.sidecar.unlink()
        self.legacy([["demo.mp4", "D:/old/demo.mp4", "published", "待审核", "人工备注"]])
        result = self.scan()
        self.assertEqual(result.items[0]["status"], "skipped")
        self.assertEqual(self.db.records()[0]["status"], "legacy_uploaded")

    def test_renamed_completed_content_is_not_reuploaded(self):
        result = self.scan()
        sha = next(iter(result.jobs))
        self.db.update(sha, status="submitted")
        (self.directory / "demo.mp4").rename(self.directory / "renamed.mp4")
        result = self.scan()
        self.assertFalse(result.jobs)
        self.assertEqual(result.items[0]["status"], "skipped")

    def test_same_content_is_only_queued_once(self):
        other = make_materials(self.directory, name="copy", episode=2)
        video = self.directory / "copy.mp4"
        video.write_bytes((self.directory / "demo.mp4").read_bytes())
        data = json.loads(other.read_text(encoding="utf-8"))
        data["mediaInfo"]["sha256"] = file_hash(video)
        atomic_json(other, data)
        result = self.scan()
        self.assertEqual(len(result.jobs), 1)
        self.assertEqual(sum(i["status"] == "duplicate" for i in result.items), 1)

    def test_valid_copy_is_not_lost_if_first_copy_lacks_materials(self):
        other = make_materials(self.directory, name="zzz")
        (self.directory / "zzz.mp4").write_bytes((self.directory / "demo.mp4").read_bytes())
        data = json.loads(other.read_text(encoding="utf-8"))
        data["mediaInfo"]["sha256"] = file_hash(self.directory / "zzz.mp4")
        atomic_json(other, data)
        self.sidecar.unlink()
        result = self.scan()
        self.assertEqual(len(result.jobs), 1)

    def test_missing_materials_block_but_testmode_can_prepare(self):
        make_materials(self.directory, name="test-mode", episode=2, test_mode=True)
        (self.directory / "orphan.mp4").write_bytes(b"orphan")
        result = self.scan()
        self.assertEqual(len(result.jobs), 2)
        reasons = "\n".join(item.get("reason", "") for item in result.items)
        self.assertTrue(any(job.manifest.testMode for job in result.jobs.values()))
        self.assertIn("缺少", reasons)

    def test_sidecar_video_path_can_be_relocated_with_verified_hash(self):
        self.modify(lambda d: d["episodes"][0].update(sourceFile="D:/no-longer-present/demo.mp4"))
        job = next(iter(self.scan().jobs.values()))
        self.assertEqual(job.files["video"], self.directory / "demo.mp4")

    def test_wrong_hash_is_rejected(self):
        self.modify(lambda d: d["mediaInfo"].update(sha256="0" * 64))
        result = self.scan()
        self.assertFalse(result.jobs)
        self.assertIn("SHA256", result.items[0]["reason"])

    def test_changed_business_data_after_platform_action_stops_resume(self):
        result = self.scan()
        sha = next(iter(result.jobs))
        self.db.update(sha, content_id="C1", status="ready")
        self.modify(lambda d: d.update(synopsis="changed"))
        self.assertFalse(self.scan().jobs)
        self.assertEqual(self.db.get(sha)["status"], "blocked")

    def test_invalid_material_does_not_erase_unknown_submit_state(self):
        sha = next(iter(self.scan().jobs))
        self.db.update(sha, status="unknown", details={"submissionAttempted": True})
        self.modify(lambda d: d.update(testMode=True))
        self.scan()
        self.assertEqual(self.db.get(sha)["status"], "unknown")

    def test_previous_version_submitted_state_is_imported(self):
        job = self.job()
        atomic_json(self.settings.root / "live" / job.job_id / "run_state.json",
                    {"fingerprint": job.fingerprint, "status": "submitted", "contentId": "C1"})
        result = self.scan()
        self.assertFalse(result.jobs)
        self.assertEqual(result.items[0]["status"], "skipped")

    def test_series_mapping_reused_in_episode_order(self):
        make_materials(self.directory, name="episode2", episode=2)
        result = self.scan()
        first = result.ordered_jobs()[0]
        self.db.remember_series(self.db.get(first[0])["series_key"], "C1")
        result = self.scan()
        self.assertEqual([job.manifest.episodes[0].number for _, job in result.ordered_jobs()], [1, 2])
        self.assertTrue(all(self.db.get(sha)["content_id"] == "C1" for sha in result.jobs))

    def test_conflicting_series_ids_are_rejected(self):
        sha = next(iter(self.scan().jobs))
        self.db.remember_series(self.db.get(sha)["series_key"], "C1")
        with self.assertRaises(ValueError):
            self.db.remember_series(self.db.get(sha)["series_key"], "C2")

    def test_ledger_preserves_notes_custom_sheets_and_creates_backup(self):
        self.legacy([["demo.mp4", str(self.directory / "demo.mp4"), "published", "待审核", "人工备注"]])
        self.scan()
        self.assertEqual(sync_ledger(self.db, Ledger(self.settings.ledger)), [])
        book = load_workbook(self.settings.ledger)
        try:
            sheet = book["发布记录"]
            self.assertEqual(sheet.max_row, 2)
            self.assertIn("人工备注", sheet["E2"].value)
            self.assertEqual(book["保留数据"]["A1"].value, "keep")
        finally:
            book.close()
        self.assertTrue(self.settings.ledger.with_name("发布记录.before-agent.xlsx").exists())

    def test_ledger_failure_keeps_completed_hash_and_pending_export(self):
        sha = next(iter(self.scan().jobs))
        self.db.update(sha, status="submitted", ledger_dirty=1)
        with patch.object(Ledger, "write", side_effect=PermissionError("Excel 占用")):
            self.assertTrue(sync_ledger(self.db, Ledger(self.settings.ledger)))
        self.assertEqual(self.db.get(sha)["status"], "submitted")
        self.assertEqual(self.db.get(sha)["ledger_dirty"], 1)
        self.assertFalse(self.scan().jobs)

    def test_duplicate_legacy_rows_fail_closed(self):
        self.legacy([["demo.mp4", str(self.directory / "demo.mp4"), "published", None, None]] * 2)
        result = self.scan()
        self.assertFalse(result.jobs)
        self.assertIn("多条", result.items[0]["reason"])

    def test_preview_does_not_modify_sqlite_or_excel(self):
        self.scan()
        before = self.db.records()
        self.legacy([["demo.mp4", str(self.directory / "demo.mp4"), "published", None, None]])
        original = self.settings.ledger.read_bytes()
        preview = UploadDatabase(self.db.path, preview=True)
        try:
            scan_directory(self.settings, preview)
        finally:
            preview.close()
        self.assertEqual(self.db.records(), before)
        self.assertEqual(self.settings.ledger.read_bytes(), original)

    def test_plain_main_dry_run_does_not_start_agent(self):
        import main
        with patch.object(main.Settings, "from_env", return_value=self.settings), patch.object(main, "process_queue", side_effect=AssertionError("不应启动")):
            self.assertEqual(main.main(["--dry-run"]), 0)
        self.assertTrue((self.settings.root / "last_scan.json").exists())

    def test_locks_refuse_duplicate_process(self):
        lock = self.settings.root / "run.lock"
        with exclusive_lock(lock):
            with self.assertRaises(RuntimeError):
                with exclusive_lock(lock):
                    pass
        self.assertFalse(lock.exists())

    def test_tools_require_duplicate_check_before_writes(self):
        context = self.context("ready")
        with self.assertRaisesRegex(ValueError, "查重"):
            asyncio.run(context.submit_upload(JobParams(job_id=context.job_id)))
        self.assertEqual(context.browser.submit.await_count, 0)

    def test_submit_intent_saved_before_browser_and_no_double_submit(self):
        context = self.context("ready", formVerified=True, cardCount=1)
        context.checked = True
        async def submit(state):
            self.assertEqual(context.record["status"], "submitting")
            self.assertTrue(context.state["submissionAttempted"])
            return {"contentId": "C1", "episodeId": "E1", "reviewStatus": "待审核"}
        context.browser.submit = AsyncMock(side_effect=submit)
        asyncio.run(context.submit_upload(JobParams(job_id=context.job_id)))
        self.assertEqual(context.record["status"], "submitted")
        with self.assertRaises(ValueError):
            asyncio.run(context.submit_upload(JobParams(job_id=context.job_id)))
        self.assertEqual(context.browser.submit.await_count, 1)

    def test_mock_upload_finishes_preparation_without_publication_or_repeat(self):
        self.modify(lambda d: d.update(testMode=True))
        context = self.context("uploading", formVerified=True, cardCount=1, uploadedAssets=["video", "cover", "copyright"])
        context.checked = True
        context.browser.snap["uploadRows"] = [{"text": "demo.mp4 100% 上传完成", "previews": 1}]
        result = asyncio.run(context.get_upload_status(JobParams(job_id=context.job_id)))
        self.assertEqual(result["status"], "test_prepared")
        self.assertFalse(result["submission_allowed"])
        self.assertNotIn("submissionAttempted", context.state)
        self.assertNotIn("submittedAt", context.state)
        context.browser.submit.assert_not_called()
        self.assertEqual(sync_ledger(self.db, Ledger(self.settings.ledger)), [])
        self.assertEqual(Ledger(self.settings.ledger).read()[0]["发布状态"], "test_prepared")
        resumed = self.scan()
        self.assertFalse(resumed.jobs)
        self.assertEqual(resumed.items[0]["status"], "test_prepared")

    def test_mock_submit_tool_is_idempotent_and_does_not_click_submit(self):
        self.modify(lambda d: d.update(testMode=True))
        context = self.context("ready", formVerified=True, cardCount=1)
        context.checked = True
        params = JobParams(job_id=context.job_id)
        self.assertEqual(asyncio.run(context.submit_upload(params))["status"], "test_prepared")
        self.assertEqual(asyncio.run(context.submit_upload(params))["status"], "test_prepared")
        context.browser.submit.assert_not_called()
        self.assertNotIn("submissionAttempted", context.state)

    def test_browser_adapter_also_refuses_final_submit_for_mock_data(self):
        self.modify(lambda d: d.update(testMode=True))
        browser = PlatformBrowser(self.job(), self.settings.root)
        browser.verify_prepared = AsyncMock()
        with self.assertRaisesRegex(ValueError, "模拟模式"):
            asyncio.run(browser.submit({}))
        browser.verify_prepared.assert_not_called()

    def test_mock_queue_starts_browser_and_accepts_prepared_terminal_state(self):
        import main
        self.modify(lambda d: d.update(testMode=True))
        report = self.scan()
        contexts = []
        def create_context(*args):
            browser = FakeBrowser(args[0])
            browser.browser = object()
            browser.open = AsyncMock()
            browser.close = AsyncMock()
            context = UploadTask(*args, browser=browser)
            context.bootstrap = AsyncMock()
            contexts.append(context)
            return context
        database = self.db
        class FakeAgent:
            def __init__(self, **kwargs):
                self.tools = kwargs["tools"]
                assert "test_prepared" in kwargs["task"]
            async def run(self, **kwargs):
                context = contexts[-1]
                context.checked = True
                database.update(context.sha256, status="ready", content_id="C1", details={"formVerified": True, "cardCount": 1})
                return await self.tools.registry.execute_action("submit_upload", {"job_id": context.job_id})
        with patch.dict(os.environ, {"OPENAI_MODEL": "deepseek-flash-free", "OPENAI_API_KEY": "local-test-only", "OPENAI_BASE_URL": "https://unused.invalid/v1",
                                     "ANONYMIZED_TELEMETRY": "false", "BROWSER_USE_CLOUD_SYNC": "false",
                                     "BROWSER_USE_CONFIG_DIR": str(self.settings.root / "browser-config")}), \
             patch("browser_use.Agent", FakeAgent), patch("browser_use.ChatDeepSeek", return_value=SimpleNamespace(model="deepseek-flash-free", provider="deepseek", name="deepseek-flash-free")), \
             patch("kuaishou_upload.tools.UploadTask", side_effect=create_context):
            results = asyncio.run(main.process_queue(self.settings, self.db, report))
        self.assertEqual(results[0]["status"], "test_prepared")
        contexts[0].browser.open.assert_awaited_once()
        contexts[0].browser.close.assert_awaited_once()
        contexts[0].browser.submit.assert_not_called()
        self.assertEqual(Ledger(self.settings.ledger).read()[0]["发布状态"], "test_prepared")

    def test_timeout_after_submit_only_reconciles(self):
        context = self.context("ready", formVerified=True, cardCount=1)
        context.checked = True
        context.browser.submit.side_effect = ConnectionError("断线")
        with self.assertRaises(ConnectionError):
            asyncio.run(context.submit_upload(JobParams(job_id=context.job_id)))
        asyncio.run(context.fail(ConnectionError("断线")))
        self.assertEqual(context.record["status"], "unknown")
        resumed = UploadTask(context.job, context.sha256, self.settings, self.db, {}, browser=FakeBrowser(context.job))
        resumed.checked = True
        with self.assertRaisesRegex(ValueError, "只核对"):
            asyncio.run(resumed.submit_upload(JobParams(job_id=resumed.job_id)))
        self.assertEqual(resumed.browser.submit.await_count, 0)

    def test_attachment_selection_intent_prevents_duplicate_upload(self):
        context = self.context("uploading", attachmentIntent="video")
        context.checked = True
        with self.assertRaisesRegex(ValueError, "选择结果不明确"):
            asyncio.run(context.upload_files(UploadParams(job_id=context.job_id, files=[{"role": "video", "ref": "k1"}])))
        self.assertEqual(context.browser.act.await_count, 0)

    def test_search_cannot_edit_series_title_in_detail_page(self):
        context = self.context()
        context.browser.snap["url"] = "https://kdj.kuaishou.com/home/content/content-management/detail?miniSeriesId=C1"
        with self.assertRaisesRegex(ValueError, "搜索框"):
            asyncio.run(context.inspect_platform(InspectParams(job_id=context.job_id, action="search", ref="k1")))
        self.assertEqual(context.browser.act.await_count, 0)

    def test_unsafe_generic_tools_not_available(self):
        context = self.context()
        with patch.dict(os.environ, {"BROWSER_USE_CONFIG_DIR": str(self.settings.root / "config"), "ANONYMIZED_TELEMETRY": "false"}):
            tools = build_tools(context)
        allowed = {"done", "scan_directory", "inspect_platform", "fill_form", "upload_files", "submit_upload", "get_upload_status"}
        self.assertEqual(set(tools.registry.registry.actions), allowed)
        result = asyncio.run(tools.registry.execute_action("scan_directory", {}))
        self.assertIn(context.job_id, result.extracted_content)

    def test_browser_use_agent_accepts_tools_without_restoring_generic_actions(self):
        context = self.context()
        with patch.dict(os.environ, {"BROWSER_USE_CONFIG_DIR": str(self.settings.root / "config"), "ANONYMIZED_TELEMETRY": "false", "BROWSER_USE_CLOUD_SYNC": "false"}):
            from browser_use import Agent, Browser, ChatDeepSeek
            tools = build_tools(context)
            agent = Agent(task="仅验证本地工具注册，不运行网页任务", llm=ChatDeepSeek(model="deepseek-flash-free", api_key="local-test-only"),
                          browser=Browser(headless=True, enable_default_extensions=False), tools=tools, use_vision=False, use_judge=False)
            names = set(agent.tools.registry.registry.actions)
            self.assertNotIn("click", names)
            self.assertNotIn("evaluate", names)
            self.assertIn("submit_upload", names)
            self.assertIn("scan_directory", names)
            asyncio.run(agent.close())

    def test_buttons_cannot_bypass_submit_tool(self):
        for text in ("提交", "发布", "立即发布", "确认发布", "删除", "撤回"):
            self.assertFalse(safe_control({"text": text}, "prepare", {"剧名": text}, []))

    def test_video_success_cannot_use_another_file_progress(self):
        self.assertFalse(PlatformBrowser.video_complete({"uploadRows": [{"text": "demo.mp4 0%", "previews": 1},
                                                                                     {"text": "proof.png 100%", "previews": 1}]}, "demo.mp4"))

    def test_boolean_authorization_and_legacy_episode_key(self):
        self.modify(lambda d: d["episodes"][0].update(no=d["episodes"][0].pop("number")))
        self.assertEqual(self.job().manifest.episodes[0].number, 1)
        self.modify(lambda d: d["copyright"].update(authorizedForUpload="false"))
        with self.assertRaises(ValueError):
            self.job()


if __name__ == "__main__":
    unittest.main()
